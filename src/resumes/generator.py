import os
import json
import re
import datetime
import google.generativeai as genai
import docx
from docx import Document
from playwright.sync_api import sync_playwright
from src.vault.parser import parse_markdown_file, write_markdown_file

class EvidenceAgent:
    def __init__(self, vault_path: str = "vault"):
        self.vault_path = vault_path
        api_key = os.environ.get("GEMINI_API_KEY")
        if api_key:
            genai.configure(api_key=api_key)

    def load_candidate_knowledge_base(self) -> dict:
        kb = {"experience": [], "skills": [], "projects": []}
        
        # Load Experience
        exp_dir = os.path.join(self.vault_path, "05 - Candidate Profile", "Experience")
        if os.path.exists(exp_dir):
            for f in os.listdir(exp_dir):
                if f.endswith(".md"):
                    fm, body = parse_markdown_file(os.path.join(exp_dir, f))
                    kb["experience"].append({"file": f, "fm": fm, "body": body})
                    
        # Load Skills
        skills_dir = os.path.join(self.vault_path, "05 - Candidate Profile", "Skills")
        if os.path.exists(skills_dir):
            for f in os.listdir(skills_dir):
                if f.endswith(".md"):
                    fm, body = parse_markdown_file(os.path.join(skills_dir, f))
                    kb["skills"].append({"file": f, "fm": fm, "body": body})
                    
        # Load Projects
        proj_dir = os.path.join(self.vault_path, "05 - Candidate Profile", "Projects")
        if os.path.exists(proj_dir):
            for f in os.listdir(proj_dir):
                if f.endswith(".md"):
                    fm, body = parse_markdown_file(os.path.join(proj_dir, f))
                    kb["projects"].append({"file": f, "fm": fm, "body": body})
                    
        return kb

    def map_evidence(self, job_filepath: str) -> str:
        job_fm, job_body = parse_markdown_file(job_filepath)
        job_id = job_fm.get("job_id")
        
        kb = self.load_candidate_knowledge_base()
        
        # Call LLM to map job requirements to candidate evidence
        prompt = f"""
Given the following Candidate Knowledge Base (experience, skills, projects) and the Job Description,
map each key job requirement to candidate evidence. 

For each requirement, classify it as: STRONG_MATCH, MATCH, PARTIAL_MATCH, TRANSFERABLE, MISSING, or UNKNOWN.
Use Obsidian wiki-links (e.g. [[SQL]], [[Company A]], [[Healthcare MDM]]) to point to the source of evidence.

Candidate Knowledge Base:
---
Experience:
{json.dumps(kb["experience"], indent=2)}

Skills:
{json.dumps(kb["skills"], indent=2)}

Projects:
{json.dumps(kb["projects"], indent=2)}
---

Job Details:
Title: {job_fm.get('title')}
Company: {job_fm.get('company')}
Job Content:
{job_body}

Return ONLY raw JSON in this format:
{{
  "requirements": [
    {{
      "requirement_name": "SQL",
      "requirement_description": "Strong SQL skills",
      "evidence": ["[[SQL]]", "[[Company A]]", "[[Healthcare MDM]]"],
      "classification": "STRONG_MATCH"
    }}
  ]
}}
Do not wrap it in markdown code blocks.
"""
        model = genai.GenerativeModel('gemini-1.5-flash')
        response = model.generate_content(prompt)
        text = response.text.strip()
        if text.startswith("```"):
            text = re.sub(r"^```[a-zA-Z]*\n", "", text)
            text = re.sub(r"\n```$", "", text)
            text = text.strip()
            
        map_data = json.loads(text)
        
        # Write to vault/02 - Job Analysis/<JOB_ID> - Evidence Map.md
        analysis_dir = os.path.join(self.vault_path, "02 - Job Analysis")
        os.makedirs(analysis_dir, exist_ok=True)
        filepath = os.path.join(analysis_dir, f"{job_id} - Evidence Map.md")
        
        fm = {
            "type": "evidence_map",
            "job_id": job_id
        }
        
        body = "# Evidence Map\n\n"
        for req in map_data.get("requirements", []):
            body += f"## {req['requirement_name']}\n\n"
            body += f"JD Requirement:\n\n\"{req['requirement_description']}\"\n\n"
            body += "Candidate Evidence:\n\n"
            body += "\n".join([f"- {ev}" for ev in req['evidence']]) + "\n\n"
            body += f"Classification:\n\n{req['classification']}\n\n"
            
        write_markdown_file(filepath, fm, body)
        return filepath


class ResumeAgent:
    def __init__(self, vault_path: str = "vault"):
        self.vault_path = vault_path
        api_key = os.environ.get("GEMINI_API_KEY")
        if api_key:
            genai.configure(api_key=api_key)

    def generate_resume(self, job_filepath: str, evidence_map_filepath: str) -> str:
        job_fm, job_body = parse_markdown_file(job_filepath)
        job_id = job_fm.get("job_id")
        
        _, evidence_body = parse_markdown_file(evidence_map_filepath)
        
        # Load profile
        profile_path = os.path.join(self.vault_path, "05 - Candidate Profile/Profile.md")
        profile_fm, profile_body = parse_markdown_file(profile_path)
        
        prompt = f"""
You are an expert resume writer. Tailor a resume for the candidate based ONLY on the evidence provided in the Evidence Map and Candidate Profile.
Strictly adhere to the following rules:
1. NEVER fabricate any skills, metrics, certifications, or experience.
2. If evidence does not exist for a requirement, omit it or classify as MISSING or UNKNOWN. Do not upgrade to match.
3. Optimize representation: prioritize relevant experience, reorder skills, and highlight relevant projects.
4. Keep the chronology and actual job titles factual and unaltered.

Candidate Profile:
---
{profile_body}
---

Evidence Map:
---
{evidence_body}
---

Job Title: {job_fm.get('title')}
Company: {job_fm.get('company')}

Generate a clean markdown resume. Do not include HTML. Output only the markdown resume.
"""
        model = genai.GenerativeModel('gemini-1.5-flash')
        response = model.generate_content(prompt)
        resume_md = response.text.strip()
        
        tailored_dir = os.path.join(self.vault_path, "04 - Resumes", "Tailored", job_id)
        os.makedirs(tailored_dir, exist_ok=True)
        
        resume_path = os.path.join(tailored_dir, "03 - Resume.md")
        fm = {
            "type": "resume",
            "job_id": job_id,
            "version": "1.0",
            "generated_date": datetime.date.today().isoformat()
        }
        write_markdown_file(resume_path, fm, resume_md)
        
        # Write docx
        self.convert_md_to_docx(resume_md, os.path.join(tailored_dir, "Resume.docx"))
        
        # Write pdf
        self.convert_md_to_pdf(resume_md, os.path.join(tailored_dir, "Resume.pdf"))
        
        return resume_path

    def convert_md_to_docx(self, markdown_text: str, output_path: str):
        doc = Document()
        for line in markdown_text.split("\n"):
            line = line.strip()
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line:
                doc.add_paragraph(line)
        doc.save(output_path)

    def convert_md_to_pdf(self, markdown_text: str, output_path: str):
        # We can construct HTML from MD and print using Playwright
        import markdown
        html_content = markdown.markdown(markdown_text)
        
        styled_html = f"""
        <html>
        <head>
            <style>
                body {{
                    font-family: Arial, sans-serif;
                    line-height: 1.5;
                    margin: 40px;
                    color: #333;
                }}
                h1 {{ text-align: center; border-bottom: 2px solid #333; padding-bottom: 10px; }}
                h2 {{ border-bottom: 1px solid #666; padding-bottom: 5px; margin-top: 20px; }}
                ul {{ margin-left: 20px; }}
            </style>
        </head>
        <body>
            {html_content}
        </body>
        </html>
        """
        
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page()
            page.set_content(styled_html)
            page.pdf(path=output_path, format="A4")
            browser.close()


class ATSValidationAgent:
    def __init__(self, vault_path: str = "vault"):
        self.vault_path = vault_path
        api_key = os.environ.get("GEMINI_API_KEY")
        if api_key:
            genai.configure(api_key=api_key)

    def validate_resume(self, job_filepath: str, resume_filepath: str, evidence_map_filepath: str) -> dict:
        job_fm, job_body = parse_markdown_file(job_filepath)
        _, resume_body = parse_markdown_file(resume_filepath)
        _, evidence_body = parse_markdown_file(evidence_map_filepath)
        job_id = job_fm.get("job_id")
        
        prompt = f"""
Perform an ATS audit comparing the target Job Description, Candidate Resume, and Evidence Map.
Verify that:
1. No unsupported or fabricated candidate claims exist in the resume.
2. The resume matches the terminology and structure of the job description.

Job Description:
---
{job_body}
---

Candidate Resume:
---
{resume_body}
---

Evidence Map:
---
{evidence_body}
---

Return a JSON response containing:
- application_match_score (integer 0-100)
- strong_coverage (list of strings)
- partial_coverage (list of strings)
- missing (list of strings)
- unsupported_claims_detected (boolean)
- revision_needed (boolean, set to true if score < 85 or if unsupported claims were detected)
- revision_suggestions (string)

Return ONLY raw JSON, starting with {{ and ending with }}. Do not wrap it in markdown code blocks.
"""
        model = genai.GenerativeModel('gemini-1.5-flash')
        response = model.generate_content(prompt)
        text = response.text.strip()
        if text.startswith("```"):
            text = re.sub(r"^```[a-zA-Z]*\n", "", text)
            text = re.sub(r"\n```$", "", text)
            text = text.strip()
            
        res = json.loads(text)
        
        # Save report
        tailored_dir = os.path.join(self.vault_path, "04 - Resumes", "Tailored", job_id)
        report_path = os.path.join(tailored_dir, "04 - ATS Validation.md")
        
        fm = {
            "type": "ats_validation",
            "job_id": job_id,
            "application_match_score": res["application_match_score"]
        }
        
        body = f"""# ATS / Application Validation

## Application Match Score

{res['application_match_score']}/100

## Strong Coverage

{chr(10).join(['- ' + sc for sc in res.get('strong_coverage', [])])}

## Partial Coverage

{chr(10).join(['- ' + pc for pc in res.get('partial_coverage', [])])}

## Missing

{chr(10).join(['- ' + m for m in res.get('missing', [])])}

## Evidence Integrity

{"FAIL - Unsupported candidate claims detected!" if res.get('unsupported_claims_detected') else "PASS - No unsupported candidate claims detected."}

## Suggestions

{res.get('revision_suggestions', 'None.')}
"""
        write_markdown_file(report_path, fm, body)
        return res
